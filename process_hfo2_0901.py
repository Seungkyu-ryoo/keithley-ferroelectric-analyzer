#!/usr/bin/env python3
"""Build processed HfO2 workbooks and Pubfig projects for 2026-09-01.

The numerical transformations intentionally mirror
``ferro_cycle_analyzer_fixed.py``.  The Pubfig v3 project layout and plot
styling are cloned from the supplied 2026-08-26 5 nm project, while all data,
column identifiers, graph references, and tree nodes are regenerated.
"""

from __future__ import annotations

import argparse
import contextlib
import copy
import io
import json
import math
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable
from uuid import uuid4

import numpy as np
import pandas as pd
from matplotlib import colormaps
from matplotlib.colors import to_hex
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.integrate import cumulative_trapezoid
from scipy.interpolate import interp1d


AREA_UM2 = 400.0
DATE_TOKEN = "0901"
DATE_DISPLAY = "09/01"
MODES = ("Endurance", "PUND", "PV")
MODE_ORDER = {name: index for index, name in enumerate(MODES)}
METRICS = (
    "Ec+",
    "Ec-",
    "|Ec+|",
    "|Ec-|",
    "|Ec+|-|Ec-|",
    "2Ec",
    "Pr+",
    "Pr-",
    "|Pr+|",
    "|Pr-|",
    "|Pr+|-|Pr-|",
    "2Pr",
)
METRIC_COLUMNS = ("condition", "cycle", "file", *METRICS)
MANIFEST_COLUMNS = (
    "pulse_condition",
    "condition",
    "mode",
    "cycle_file_value",
    "source_file",
    "source_relative_path",
    "input_rows",
    "output_rows",
    "status",
    "warning",
    "size_bytes",
)
ENDURANCE_COLUMNS = (
    "condition",
    "file",
    "global_cycle",
    "pundEndurance",
    "iteration",
    "P",
    "Pa",
    "U",
    "Ua",
    "N",
    "Na",
    "D",
    "Da",
    "Psw",
    "Qsw",
    "plot_cycle",
    "axis_segment",
    "axis_cycle",
    "Psw_uC_cm2",
    "Qsw_uC_cm2",
)
PUND_LOOP_COLUMNS = (
    "condition",
    "cycle",
    "file",
    "Voltage_Pos_V",
    "P_Pos_uC_cm2",
    "Voltage_Neg_V",
    "P_Neg_uC_cm2",
)
PV_LOOP_COLUMNS = (
    "condition",
    "cycle",
    "file",
    "Voltage_V",
    "Field_MV_cm",
    "P_uC_cm2",
    "J_A_cm2",
)


@dataclass(frozen=True)
class GroupSpec:
    thickness_label: str
    thickness_nm: float
    file_group: str
    relative_source: str
    json_subdir: str
    allowed_leaf_paths: tuple[str, ...] | None = None

    @property
    def basename(self) -> str:
        return f"HfO2_{self.thickness_label}_{self.file_group}"


DEFAULT_GROUPS = (
    GroupSpec("5nm", 5.0, "2.0ms", "HfO2_5nm/L7_716V_2ms", "5nm"),
    GroupSpec("10nm", 10.0, "Asdep", "HfO2_10nm/Asdep", "10nm"),
    GroupSpec("10nm", 10.0, "0.2ms", "HfO2_10nm/716V_0.2ms", "10nm"),
    GroupSpec("10nm", 10.0, "0.4ms", "HfO2_10nm/716V_0.4ms", "10nm"),
)


@dataclass
class Condition:
    condition_id: str
    display_name: str
    voltage: str
    voltage_value: float
    column_name: str | None
    leaf_relative: str
    files: dict[str, list[Path]] = field(
        default_factory=lambda: {mode: [] for mode in MODES}
    )


@dataclass
class ProcessedGroup:
    spec: GroupSpec
    source_dir: Path
    conditions: list[Condition]
    metadata: pd.DataFrame
    manifest: pd.DataFrame
    tables: dict[str, dict[str, pd.DataFrame]]


def new_id(prefix: str) -> str:
    return f"{prefix}{uuid4().hex[:8]}"


def read_excel_quiet(path: Path, sheet_name: str, header: int | None = 0) -> pd.DataFrame:
    with (
        contextlib.redirect_stdout(io.StringIO()),
        contextlib.redirect_stderr(io.StringIO()),
    ):
        return pd.read_excel(
            path,
            sheet_name=sheet_name,
            header=header,
            engine="xlrd",
        )


def get_setting(settings_df: pd.DataFrame, name: str) -> float:
    rows = settings_df[
        settings_df.iloc[:, 0].astype(str).str.strip() == name
    ]
    if rows.empty:
        raise ValueError(f"Missing setting {name!r}")
    return float(rows.iloc[0, 3])


def parse_cycle(path: Path) -> int | float:
    try:
        value = float(path.stem)
    except ValueError as exc:
        raise ValueError(f"Cannot parse cycle from {path.name!r}") from exc
    return int(value) if value.is_integer() else value


def numeric_suffix(text: str | None) -> tuple[int, str]:
    if not text:
        return (-1, "")
    match = re.search(r"-?\d+", text)
    return (int(match.group()) if match else math.inf, text)


def condition_sort_key(condition: Condition) -> tuple[float, tuple[int, str], str]:
    return (
        condition.voltage_value,
        numeric_suffix(condition.column_name),
        condition.condition_id,
    )


def discover_conditions(
    campaign_root: Path,
    spec: GroupSpec,
    *,
    date_token: str,
    date_display: str,
) -> tuple[Path, list[Condition]]:
    source_dir = campaign_root / spec.relative_source
    if not source_dir.is_dir():
        raise FileNotFoundError(f"Missing source group: {source_dir}")

    allowed = (
        {Path(item).as_posix() for item in spec.allowed_leaf_paths}
        if spec.allowed_leaf_paths
        else None
    )
    voltage_pattern = re.compile(r"^(-?\d+(?:\.\d+)?)V$", re.IGNORECASE)
    column_pattern = re.compile(r"^col\d+(?:-\d+)?$", re.IGNORECASE)
    discovered: dict[str, Condition] = {}
    leaf_by_id: dict[str, str] = {}

    for path in sorted(source_dir.rglob("*.xls")):
        mode = path.parent.name
        if mode not in MODES:
            continue
        relative = path.relative_to(source_dir)
        mode_index = len(relative.parts) - 2
        ancestors = relative.parts[:mode_index]
        leaf_relative = Path(*ancestors).as_posix()
        if allowed is not None and leaf_relative not in allowed:
            continue

        voltage_candidates = [
            (part, voltage_pattern.match(part)) for part in ancestors
        ]
        voltage_candidates = [item for item in voltage_candidates if item[1]]
        if not voltage_candidates:
            raise ValueError(f"No voltage directory found above {path}")
        voltage, voltage_match = voltage_candidates[-1]
        assert voltage_match is not None
        voltage_value = float(voltage_match.group(1))
        columns = [part for part in ancestors if column_pattern.match(part)]
        column_name = columns[-1] if columns else None

        pieces = [voltage]
        if column_name:
            pieces.append(column_name)
        pieces.append(date_token)
        condition_id = "_".join(pieces)
        display_parts = [voltage]
        if column_name:
            display_parts.append(column_name)
        display_name = f"{' '.join(display_parts)} ({date_display})"

        previous_leaf = leaf_by_id.get(condition_id)
        if previous_leaf is not None and previous_leaf != leaf_relative:
            raise ValueError(
                f"Condition {condition_id!r} maps to both {previous_leaf!r} "
                f"and {leaf_relative!r}"
            )
        leaf_by_id[condition_id] = leaf_relative
        condition = discovered.setdefault(
            condition_id,
            Condition(
                condition_id=condition_id,
                display_name=display_name,
                voltage=voltage,
                voltage_value=voltage_value,
                column_name=column_name,
                leaf_relative=leaf_relative,
            ),
        )
        condition.files[mode].append(path)

    if not discovered:
        raise ValueError(f"No Keithley XLS files found in {source_dir}")

    for condition in discovered.values():
        for mode in MODES:
            paths = condition.files[mode]
            paths.sort(key=lambda item: (float(parse_cycle(item)), item.name))
            seen: set[int | float] = set()
            for path in paths:
                cycle = parse_cycle(path)
                if cycle in seen:
                    raise ValueError(
                        f"Duplicate {mode} cycle {cycle:g} for "
                        f"{condition.condition_id}"
                    )
                seen.add(cycle)

    return source_dir, sorted(discovered.values(), key=condition_sort_key)


def charge_to_polarization(values: Any, area_um2: float = AREA_UM2) -> Any:
    return values / (area_um2 * 1e-8) * 1e6


def current_to_density(values: Any, area_um2: float = AREA_UM2) -> Any:
    return values / (area_um2 * 1e-8)


def voltage_to_field(values: Any, thickness_nm: float) -> Any:
    return values * 10.0 / thickness_nm


def calculate_pund_file(path: Path) -> tuple[pd.DataFrame, int]:
    data_df = read_excel_quiet(path, "Data")
    settings_df = read_excel_quiet(path, "Settings", header=None)
    needed = {"t", "V", "I"}
    if not needed.issubset(data_df.columns):
        missing = sorted(needed.difference(data_df.columns))
        raise ValueError(f"Missing PUND columns: {missing}")

    times = data_df["t"].to_numpy(dtype=float)
    volts = data_df["V"].to_numpy(dtype=float)
    currents = data_df["I"].to_numpy(dtype=float)
    tp = get_setting(settings_df, "tp")
    td = get_setting(settings_df, "td")
    trf = get_setting(settings_df, "trf")
    pulse_duration = trf + tp + trf
    pulse_interval = pulse_duration + td
    starts = [td + index * pulse_interval for index in range(4)]
    t_common = np.linspace(0.0, pulse_duration, 2000)

    def interp_pulse(start: float) -> tuple[np.ndarray, np.ndarray]:
        mask = (times >= start) & (times <= start + pulse_duration)
        if int(mask.sum()) < 2:
            raise ValueError(
                f"Insufficient PUND samples in pulse starting at {start:g}"
            )
        t_relative = times[mask] - start
        return (
            interp1d(
                t_relative,
                volts[mask],
                bounds_error=False,
                fill_value=0,
            )(t_common),
            interp1d(
                t_relative,
                currents[mask],
                bounds_error=False,
                fill_value=0,
            )(t_common),
        )

    voltage_p, current_p = interp_pulse(starts[0])
    _, current_u = interp_pulse(starts[1])
    voltage_n, current_n = interp_pulse(starts[2])
    _, current_d = interp_pulse(starts[3])
    charge_pos = cumulative_trapezoid(
        current_p - current_u, t_common, initial=0
    )
    charge_neg = cumulative_trapezoid(
        current_n - current_d, t_common, initial=0
    )
    polarization_pos = charge_to_polarization(charge_pos)
    polarization_neg = charge_to_polarization(charge_neg)
    polarization_pos -= (polarization_pos[0] + polarization_pos[-1]) / 2.0
    polarization_neg -= (polarization_neg[0] + polarization_neg[-1]) / 2.0
    result = pd.DataFrame(
        {
            "Voltage_Pos_V": voltage_p,
            "P_Pos_uC_cm2": polarization_pos,
            "Voltage_Neg_V": voltage_n,
            "P_Neg_uC_cm2": polarization_neg,
        }
    )
    return result, len(data_df)


def crossing_values(x_data: Any, y_data: Any, target: float) -> list[float]:
    x_values = np.asarray(x_data, dtype=float)
    y_values = np.asarray(y_data, dtype=float)
    values: list[float] = []
    for index in range(len(x_values) - 1):
        x0, x1 = x_values[index], x_values[index + 1]
        y0, y1 = y_values[index], y_values[index + 1]
        if not np.isfinite([x0, x1, y0, y1]).all():
            continue
        delta0 = y0 - target
        delta1 = y1 - target
        if delta0 == 0:
            values.append(float(x0))
        elif delta0 * delta1 < 0 and y1 != y0:
            values.append(float(x0 + (target - y0) * (x1 - x0) / (y1 - y0)))
    return values


def interpolate_at_x(x_data: Any, y_data: Any, target: float) -> list[float]:
    x_values = np.asarray(x_data, dtype=float)
    y_values = np.asarray(y_data, dtype=float)
    values: list[float] = []
    for index in range(len(x_values) - 1):
        x0, x1 = x_values[index], x_values[index + 1]
        y0, y1 = y_values[index], y_values[index + 1]
        if not np.isfinite([x0, x1, y0, y1]).all():
            continue
        delta0 = x0 - target
        delta1 = x1 - target
        if delta0 == 0:
            values.append(float(y0))
        elif delta0 * delta1 < 0 and x1 != x0:
            values.append(float(y0 + (target - x0) * (y1 - y0) / (x1 - x0)))
    return values


def metric_values(
    cycle: int | float,
    file_name: str,
    ec_pos: float,
    ec_neg: float,
    pr_pos: float,
    pr_neg: float,
) -> dict[str, Any]:
    ec_finite = np.isfinite([ec_pos, ec_neg]).all()
    pr_finite = np.isfinite([pr_pos, pr_neg]).all()
    return {
        "cycle": cycle,
        "file": file_name,
        "Ec+": ec_pos,
        "Ec-": ec_neg,
        "|Ec+|": abs(ec_pos) if np.isfinite(ec_pos) else np.nan,
        "|Ec-|": abs(ec_neg) if np.isfinite(ec_neg) else np.nan,
        "|Ec+|-|Ec-|": abs(ec_pos) - abs(ec_neg) if ec_finite else np.nan,
        "2Ec": abs(ec_pos) + abs(ec_neg) if ec_finite else np.nan,
        "Pr+": pr_pos,
        "Pr-": pr_neg,
        "|Pr+|": abs(pr_pos) if np.isfinite(pr_pos) else np.nan,
        "|Pr-|": abs(pr_neg) if np.isfinite(pr_neg) else np.nan,
        "|Pr+|-|Pr-|": abs(pr_pos) - abs(pr_neg) if pr_finite else np.nan,
        "2Pr": pr_pos - pr_neg if pr_finite else np.nan,
    }


def pund_metric_row(
    loop_df: pd.DataFrame,
    cycle: int | float,
    file_name: str,
) -> dict[str, Any]:
    pos_voltage = loop_df["Voltage_Pos_V"].to_numpy(dtype=float)
    pos_pol = loop_df["P_Pos_uC_cm2"].to_numpy(dtype=float)
    neg_voltage = loop_df["Voltage_Neg_V"].to_numpy(dtype=float)
    neg_pol = loop_df["P_Neg_uC_cm2"].to_numpy(dtype=float)
    ec_pos_candidates = [
        value
        for value in crossing_values(pos_voltage, pos_pol, 0.0)
        if value > 0
    ]
    ec_neg_candidates = [
        value
        for value in crossing_values(neg_voltage, neg_pol, 0.0)
        if value < 0
    ]
    ec_pos = min(ec_pos_candidates, key=abs) if ec_pos_candidates else np.nan
    ec_neg = min(ec_neg_candidates, key=abs) if ec_neg_candidates else np.nan

    pos_start = int(np.nanargmax(pos_voltage)) if len(pos_voltage) else 0
    neg_start = int(np.nanargmin(neg_voltage)) if len(neg_voltage) else 0
    pr_pos = (
        pos_pol[pos_start + int(np.nanargmin(np.abs(pos_voltage[pos_start:])))]
        if len(pos_voltage[pos_start:])
        else np.nan
    )
    pr_neg = (
        neg_pol[neg_start + int(np.nanargmin(np.abs(neg_voltage[neg_start:])))]
        if len(neg_voltage[neg_start:])
        else np.nan
    )
    return metric_values(cycle, file_name, ec_pos, ec_neg, pr_pos, pr_neg)


def pv_offset(voltage: pd.Series, polarization: pd.Series) -> float:
    minus_pr = polarization.iloc[0]
    plus_pr: float | None = None
    threshold = 0
    for index in range(len(voltage)):
        if voltage.iloc[index] > 0.2:
            threshold = index
            break
    for index in range(threshold, len(voltage)):
        if voltage.iloc[index] < 0:
            plus_pr = float(polarization.iloc[index])
            break
    return (
        float(polarization.median())
        if plus_pr is None
        else (float(minus_pr) + plus_pr) / 2.0
    )


def calculate_pv_file(
    path: Path,
    cycle: int | float,
    thickness_nm: float,
) -> tuple[pd.DataFrame, dict[str, Any], int]:
    data_df = read_excel_quiet(path, "Data")
    needed = {"Vforce", "Charge", "Imeas"}
    if not needed.issubset(data_df.columns):
        missing = sorted(needed.difference(data_df.columns))
        raise ValueError(f"Missing PV columns: {missing}")
    voltage = pd.to_numeric(data_df["Vforce"], errors="coerce")
    polarization = charge_to_polarization(
        pd.to_numeric(data_df["Charge"], errors="coerce")
    )
    polarization = polarization - pv_offset(voltage, polarization)
    current_density = current_to_density(
        pd.to_numeric(data_df["Imeas"], errors="coerce")
    )
    loop_df = pd.DataFrame(
        {
            "Voltage_V": voltage,
            "Field_MV_cm": voltage_to_field(voltage, thickness_nm),
            "P_uC_cm2": polarization,
            "J_A_cm2": current_density,
        }
    )

    ec_crossings = crossing_values(voltage, polarization, 0.0)
    positive_ec = [value for value in ec_crossings if value > 0]
    negative_ec = [value for value in ec_crossings if value < 0]
    ec_pos = min(positive_ec, key=abs) if positive_ec else np.nan
    ec_neg = min(negative_ec, key=abs) if negative_ec else np.nan
    pr_crossings = interpolate_at_x(voltage, polarization, 0.0)
    positive_pr = [value for value in pr_crossings if value > 0]
    negative_pr = [value for value in pr_crossings if value < 0]
    pr_pos = max(positive_pr) if positive_pr else np.nan
    pr_neg = min(negative_pr) if negative_pr else np.nan
    metric = metric_values(cycle, path.name, ec_pos, ec_neg, pr_pos, pr_neg)
    return loop_df, metric, len(data_df)


def endurance_collapse_warning(block: pd.DataFrame) -> str | None:
    if len(block) < 2:
        return None
    switched = block[["Psw_uC_cm2", "Qsw_uC_cm2"]].abs()
    baseline = float(switched.iloc[:-1].stack().median())
    final_value = float(switched.iloc[-1].mean())
    if np.isfinite([baseline, final_value]).all() and baseline > 1.0:
        if final_value < baseline * 0.20:
            return "measured polarization collapse retained (possible device breakdown)"
    return None


def empty_frame(columns: Iterable[str]) -> pd.DataFrame:
    return pd.DataFrame(columns=list(columns))


def process_condition(
    condition: Condition,
    spec: GroupSpec,
    campaign_root: Path,
) -> tuple[dict[str, pd.DataFrame], list[dict[str, Any]]]:
    output = {
        "Endurance": empty_frame(ENDURANCE_COLUMNS),
        "PUND loops": empty_frame(PUND_LOOP_COLUMNS),
        "PUND metrics": empty_frame(METRIC_COLUMNS),
        "PV loops": empty_frame(PV_LOOP_COLUMNS),
        "PV metrics": empty_frame(METRIC_COLUMNS),
    }
    manifest_rows: list[dict[str, Any]] = []

    previous_end = 0
    endurance_blocks: list[pd.DataFrame] = []
    for path in condition.files["Endurance"]:
        cycle = parse_cycle(path)
        input_rows = 0
        output_rows = 0
        status = "ok"
        warning: str | None = None
        try:
            data_df = read_excel_quiet(path, "Data")
            settings_df = read_excel_quiet(path, "Settings", header=None)
            input_rows = len(data_df)
            max_loops = int(get_setting(settings_df, "max_loops"))
            missing = {"iteration", "Psw", "Qsw"}.difference(data_df.columns)
            if missing:
                raise ValueError(f"Missing Endurance columns: {sorted(missing)}")
            block = data_df.copy()
            block.insert(0, "file", path.name)
            block.insert(1, "global_cycle", previous_end + block["iteration"])
            block["plot_cycle"] = block["global_cycle"].clip(lower=1)
            block["axis_segment"] = 0
            block["axis_cycle"] = block["plot_cycle"]
            block["Psw_uC_cm2"] = charge_to_polarization(block["Psw"])
            block["Qsw_uC_cm2"] = charge_to_polarization(block["Qsw"])
            block.insert(0, "condition", condition.condition_id)
            block = block.reindex(columns=ENDURANCE_COLUMNS)
            output_rows = len(block)
            warning = endurance_collapse_warning(block)
            if warning:
                status = "warning"
            endurance_blocks.append(block)
            previous_end += max_loops
        except Exception as exc:  # retain a manifest record for every source file
            status = "error"
            warning = f"{type(exc).__name__}: {exc}"
        manifest_rows.append(
            manifest_record(
                spec,
                condition,
                "Endurance",
                cycle,
                path,
                campaign_root,
                input_rows,
                output_rows,
                status,
                warning,
            )
        )
    if endurance_blocks:
        output["Endurance"] = pd.concat(endurance_blocks, ignore_index=True)

    pund_loops: list[pd.DataFrame] = []
    pund_metrics: list[dict[str, Any]] = []
    for path in condition.files["PUND"]:
        cycle = parse_cycle(path)
        input_rows = 0
        output_rows = 0
        status = "ok"
        warning = None
        try:
            loop_df, input_rows = calculate_pund_file(path)
            output_rows = len(loop_df)
            loop_df.insert(0, "file", path.name)
            loop_df.insert(0, "cycle", cycle)
            loop_df.insert(0, "condition", condition.condition_id)
            loop_df = loop_df.reindex(columns=PUND_LOOP_COLUMNS)
            pund_loops.append(loop_df)
            metric = pund_metric_row(loop_df, cycle, path.name)
            metric["condition"] = condition.condition_id
            pund_metrics.append(metric)
        except Exception as exc:
            status = "error"
            warning = f"{type(exc).__name__}: {exc}"
        manifest_rows.append(
            manifest_record(
                spec,
                condition,
                "PUND",
                cycle,
                path,
                campaign_root,
                input_rows,
                output_rows,
                status,
                warning,
            )
        )
    if pund_loops:
        output["PUND loops"] = pd.concat(pund_loops, ignore_index=True)
    if pund_metrics:
        output["PUND metrics"] = pd.DataFrame(pund_metrics).reindex(
            columns=METRIC_COLUMNS
        )

    pv_loops: list[pd.DataFrame] = []
    pv_metrics: list[dict[str, Any]] = []
    for path in condition.files["PV"]:
        cycle = parse_cycle(path)
        input_rows = 0
        output_rows = 0
        status = "ok"
        warning = None
        try:
            loop_df, metric, input_rows = calculate_pv_file(
                path, cycle, spec.thickness_nm
            )
            output_rows = len(loop_df)
            loop_df.insert(0, "file", path.name)
            loop_df.insert(0, "cycle", cycle)
            loop_df.insert(0, "condition", condition.condition_id)
            loop_df = loop_df.reindex(columns=PV_LOOP_COLUMNS)
            pv_loops.append(loop_df)
            metric["condition"] = condition.condition_id
            pv_metrics.append(metric)
        except Exception as exc:
            status = "error"
            warning = f"{type(exc).__name__}: {exc}"
        manifest_rows.append(
            manifest_record(
                spec,
                condition,
                "PV",
                cycle,
                path,
                campaign_root,
                input_rows,
                output_rows,
                status,
                warning,
            )
        )
    if pv_loops:
        output["PV loops"] = pd.concat(pv_loops, ignore_index=True)
    if pv_metrics:
        output["PV metrics"] = pd.DataFrame(pv_metrics).reindex(
            columns=METRIC_COLUMNS
        )
    return output, manifest_rows


def manifest_record(
    spec: GroupSpec,
    condition: Condition,
    mode: str,
    cycle: int | float,
    path: Path,
    campaign_root: Path,
    input_rows: int,
    output_rows: int,
    status: str,
    warning: str | None,
) -> dict[str, Any]:
    try:
        source_relative = path.relative_to(campaign_root.parent).as_posix()
    except ValueError:
        source_relative = path.as_posix()
    return {
        "pulse_condition": spec.file_group,
        "condition": condition.condition_id,
        "mode": mode,
        "cycle_file_value": cycle,
        "source_file": path.name,
        "source_relative_path": source_relative,
        "input_rows": input_rows,
        "output_rows": output_rows,
        "status": status,
        "warning": warning,
        "size_bytes": path.stat().st_size,
    }


def build_metadata(
    spec: GroupSpec,
    source_dir: Path,
    conditions: list[Condition],
    manifest: pd.DataFrame,
    script_path: Path,
) -> pd.DataFrame:
    condition_names = "; ".join(item.condition_id for item in conditions)
    incomplete: list[str] = []
    for condition in conditions:
        missing = [mode for mode in MODES if not condition.files[mode]]
        if missing:
            incomplete.append(f"{condition.condition_id}: missing {', '.join(missing)}")
    source_modes = [
        mode
        for mode in MODES
        if any(condition.files[mode] for condition in conditions)
    ]
    warning_count = int((manifest["status"] == "warning").sum())
    error_count = int((manifest["status"] == "error").sum())
    metadata_rows: list[tuple[str, Any]] = [
        ("material", "HfO2"),
        ("nominal_thickness_nm", spec.thickness_nm),
        ("pulse_width_file_group", spec.file_group),
        ("source_pulse_directory", source_dir.name),
        (
            "pulse_width_note",
            "Pulse/flash label comes from the directory name and is not present in XLS Settings.",
        ),
        ("conditions", condition_names),
        ("source_root", str(source_dir)),
        ("source_file_count", len(manifest)),
        ("source_modes", ", ".join(source_modes)),
        ("electrode_area_um2", AREA_UM2),
        (
            "electrode_area_source",
            "ferro_cycle_analyzer_fixed.py GUI default; no area field exists in XLS workbooks.",
        ),
        ("thickness_source", f"{spec.thickness_label} source directory name"),
        ("metric_ec_axis", "Voltage (V), matching analyzer GUI default"),
        (
            "pund_curve_definition",
            "One loop per cycle: P-U (2000 points) followed by N-D (2000 points) in Pubfig JSON.",
        ),
        ("warning_file_count", warning_count),
        ("error_file_count", error_count),
        ("incomplete_conditions", "; ".join(incomplete) if incomplete else None),
        (
            "organization_note",
            "Files are grouped by flash condition and voltage; each column location is an independent device.",
        ),
        ("generated_utc", datetime.now(timezone.utc).isoformat()),
        ("analysis_script", str(script_path.resolve())),
        (
            "analysis_reference",
            str((script_path.parent / "ferro_cycle_analyzer_fixed.py").resolve()),
        ),
    ]
    return pd.DataFrame(metadata_rows, columns=["key", "value"])


def process_group(
    campaign_root: Path,
    spec: GroupSpec,
    *,
    date_token: str = DATE_TOKEN,
    date_display: str = DATE_DISPLAY,
    script_path: Path | None = None,
) -> ProcessedGroup:
    source_dir, conditions = discover_conditions(
        campaign_root,
        spec,
        date_token=date_token,
        date_display=date_display,
    )
    condition_tables: dict[str, dict[str, pd.DataFrame]] = {}
    manifest_rows: list[dict[str, Any]] = []
    for condition in conditions:
        tables, rows = process_condition(condition, spec, campaign_root)
        condition_tables[condition.condition_id] = tables
        manifest_rows.extend(rows)

    manifest = pd.DataFrame(manifest_rows).reindex(columns=MANIFEST_COLUMNS)
    if not manifest.empty:
        manifest["_mode_order"] = manifest["mode"].map(MODE_ORDER)
        manifest = (
            manifest.sort_values(
                ["condition", "_mode_order", "cycle_file_value", "source_file"],
                kind="stable",
            )
            .drop(columns="_mode_order")
            .reset_index(drop=True)
        )

    grouped_tables: dict[str, dict[str, pd.DataFrame]] = {}
    by_voltage: dict[str, list[Condition]] = defaultdict(list)
    for condition in conditions:
        by_voltage[condition.voltage].append(condition)
    for voltage in sorted(
        by_voltage,
        key=lambda item: (by_voltage[item][0].voltage_value, item),
    ):
        grouped_tables[voltage] = {}
        voltage_conditions = sorted(by_voltage[voltage], key=condition_sort_key)
        for table_name, columns in (
            ("Endurance", ENDURANCE_COLUMNS),
            ("PUND loops", PUND_LOOP_COLUMNS),
            ("PUND metrics", METRIC_COLUMNS),
            ("PV loops", PV_LOOP_COLUMNS),
            ("PV metrics", METRIC_COLUMNS),
        ):
            frames = [
                condition_tables[item.condition_id][table_name]
                for item in voltage_conditions
                if not condition_tables[item.condition_id][table_name].empty
            ]
            grouped_tables[voltage][table_name] = (
                pd.concat(frames, ignore_index=True)
                if frames
                else empty_frame(columns)
            )

    metadata = build_metadata(
        spec,
        source_dir,
        conditions,
        manifest,
        script_path or Path(__file__),
    )
    return ProcessedGroup(
        spec=spec,
        source_dir=source_dir,
        conditions=conditions,
        metadata=metadata,
        manifest=manifest,
        tables=grouped_tables,
    )


def excel_sheet_name(voltage: str, table_name: str) -> str:
    return f"{voltage}_{table_name.replace(' ', '_')}"


def write_workbook(group: ProcessedGroup, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        group.metadata.to_excel(writer, sheet_name="Metadata", index=False)
        group.manifest.to_excel(writer, sheet_name="File_manifest", index=False)
        for voltage, tables in group.tables.items():
            for table_name in (
                "Endurance",
                "PUND loops",
                "PUND metrics",
                "PV loops",
                "PV metrics",
            ):
                frame = tables[table_name]
                if frame.empty:
                    continue
                frame.to_excel(
                    writer,
                    sheet_name=excel_sheet_name(voltage, table_name),
                    index=False,
                )

        workbook = writer.book
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for column_cells in worksheet.columns:
                column_letter = column_cells[0].column_letter
                maximum = 0
                for cell in column_cells:
                    if cell.value is None:
                        continue
                    maximum = max(maximum, len(str(cell.value)))
                worksheet.column_dimensions[column_letter].width = min(
                    max(maximum + 2, 12), 55
                )


def json_scalar(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (np.bool_, bool)):
        return bool(value)
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat()
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    return value


def dataframe_records(frame: pd.DataFrame) -> list[list[Any]]:
    return [
        [json_scalar(value) for value in row]
        for row in frame.itertuples(index=False, name=None)
    ]


def table_sheet_payload(frame: pd.DataFrame) -> dict[str, Any]:
    columns = [str(column) for column in frame.columns]
    blank = ["" for _ in columns]
    return {
        "columns": columns,
        "rows": [blank, columns, *dataframe_records(frame)],
    }


def condition_token(condition_id: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]", "_", condition_id)


def metric_identifier(metric: str, token: str) -> str:
    prefix = {
        "Ec+": "Ec",
        "Ec-": "Ec-",
        "|Ec+|": "|Ec|",
        "|Ec-|": "|Ec-|",
        "|Ec+|-|Ec-|": "|Ec|-|Ec-|",
        "2Ec": "2Ec",
        "Pr+": "Pr",
        "Pr-": "Pr-",
        "|Pr+|": "|Pr|",
        "|Pr-|": "|Pr-|",
        "|Pr+|-|Pr-|": "|Pr|-|Pr-|",
        "2Pr": "2Pr",
    }[metric]
    return f"{prefix}_{token}"


def cycle_suffix(cycle: int | float) -> str:
    return f"{float(cycle):g}"


def cycle_display(cycle: int | float) -> str:
    numeric = float(cycle)
    if numeric == 1:
        return "1st cycle"
    if numeric > 0:
        exponent = math.log10(numeric)
        if abs(exponent - round(exponent)) < 1e-9:
            return f"10^{int(round(exponent))} cycle"
    return f"{numeric:g} cycle"


def wide_payload(
    columns: list[str],
    roles: list[str],
    display_names: list[str],
    values: list[list[Any]],
) -> dict[str, Any]:
    if not (len(columns) == len(roles) == len(display_names) == len(values)):
        raise ValueError("Wide-sheet metadata lengths do not match")
    maximum = max((len(items) for items in values), default=0)
    rows: list[list[Any]] = [roles, display_names]
    for row_index in range(maximum):
        rows.append(
            [
                json_scalar(items[row_index]) if row_index < len(items) else None
                for items in values
            ]
        )
    return {"columns": columns, "rows": rows}


def conditions_for_voltage(group: ProcessedGroup, voltage: str) -> list[Condition]:
    return [item for item in group.conditions if item.voltage == voltage]


def endurance_json_payload(
    group: ProcessedGroup, voltage: str
) -> tuple[dict[str, Any], list[Condition]]:
    frame = group.tables[voltage]["Endurance"]
    conditions = [
        item
        for item in conditions_for_voltage(group, voltage)
        if not frame[frame["condition"] == item.condition_id].empty
    ]
    columns: list[str] = []
    roles: list[str] = []
    labels: list[str] = []
    values: list[list[Any]] = []
    for condition in conditions:
        subset = frame[frame["condition"] == condition.condition_id]
        token = condition_token(condition.condition_id)
        columns.extend([f"cycle_{token}", f"Psw_{token}", f"Qsw_{token}"])
        roles.extend(["X", "Y", "Y"])
        labels.extend(
            [
                f"{condition.display_name} cycle",
                f"{condition.display_name} Psw",
                f"{condition.display_name} Qsw",
            ]
        )
        values.extend(
            [
                subset["plot_cycle"].tolist(),
                subset["Psw_uC_cm2"].tolist(),
                subset["Qsw_uC_cm2"].tolist(),
            ]
        )
    return wide_payload(columns, roles, labels, values), conditions


def pund_loops_json_payload(
    group: ProcessedGroup, voltage: str
) -> tuple[dict[str, Any], dict[str, list[int | float]]]:
    frame = group.tables[voltage]["PUND loops"]
    columns: list[str] = []
    roles: list[str] = []
    labels: list[str] = []
    values: list[list[Any]] = []
    cycle_map: dict[str, list[int | float]] = {}
    for condition in conditions_for_voltage(group, voltage):
        subset = frame[frame["condition"] == condition.condition_id]
        cycles = sorted(subset["cycle"].dropna().unique(), key=float)
        if not cycles:
            continue
        cycle_map[condition.condition_id] = [json_scalar(item) for item in cycles]
        token = condition_token(condition.condition_id)
        for cycle in cycles:
            cycle_rows = subset[subset["cycle"] == cycle]
            suffix = cycle_suffix(cycle)
            columns.extend([f"Vloop_{token}_{suffix}", f"Ploop_{token}_{suffix}"])
            roles.extend(["X", "Y"])
            label = cycle_display(cycle)
            labels.extend([label, label])
            voltage_values = [
                *cycle_rows["Voltage_Pos_V"].tolist(),
                *cycle_rows["Voltage_Neg_V"].tolist(),
            ]
            polarization_values = [
                *cycle_rows["P_Pos_uC_cm2"].tolist(),
                *cycle_rows["P_Neg_uC_cm2"].tolist(),
            ]
            values.extend([voltage_values, polarization_values])
    return wide_payload(columns, roles, labels, values), cycle_map


def metrics_json_payload(
    group: ProcessedGroup,
    voltage: str,
    table_name: str,
) -> tuple[dict[str, Any], list[Condition]]:
    frame = group.tables[voltage][table_name]
    columns: list[str] = []
    roles: list[str] = []
    labels: list[str] = []
    values: list[list[Any]] = []
    conditions: list[Condition] = []
    for condition in conditions_for_voltage(group, voltage):
        subset = frame[frame["condition"] == condition.condition_id].sort_values(
            "cycle"
        )
        if subset.empty:
            continue
        conditions.append(condition)
        token = condition_token(condition.condition_id)
        columns.append(f"cycle_{token}")
        roles.append("X")
        labels.append(f"{condition.display_name} cycle")
        values.append(subset["cycle"].tolist())
        for metric in METRICS:
            columns.append(metric_identifier(metric, token))
            roles.append("Y")
            labels.append(f"{condition.display_name} {metric}")
            values.append(subset[metric].tolist())
    return wide_payload(columns, roles, labels, values), conditions


def pv_loops_json_payload(
    group: ProcessedGroup, voltage: str
) -> tuple[dict[str, Any], dict[str, list[int | float]]]:
    frame = group.tables[voltage]["PV loops"]
    columns: list[str] = []
    roles: list[str] = []
    labels: list[str] = []
    values: list[list[Any]] = []
    cycle_map: dict[str, list[int | float]] = {}
    for condition in conditions_for_voltage(group, voltage):
        subset = frame[frame["condition"] == condition.condition_id]
        cycles = sorted(subset["cycle"].dropna().unique(), key=float)
        if not cycles:
            continue
        cycle_map[condition.condition_id] = [json_scalar(item) for item in cycles]
        token = condition_token(condition.condition_id)
        for cycle in cycles:
            cycle_rows = subset[subset["cycle"] == cycle]
            suffix = cycle_suffix(cycle)
            columns.extend(
                [f"V_{token}_{suffix}", f"P_{token}_{suffix}", f"J_{token}_{suffix}"]
            )
            roles.extend(["X", "Y", "Y"])
            label = cycle_display(cycle)
            labels.extend([label, label, label])
            values.extend(
                [
                    cycle_rows["Voltage_V"].tolist(),
                    cycle_rows["P_uC_cm2"].tolist(),
                    cycle_rows["J_A_cm2"].tolist(),
                ]
            )
    return wide_payload(columns, roles, labels, values), cycle_map


def extract_graph_prototypes(template: dict[str, Any]) -> dict[str, dict[str, Any]]:
    sheet_names = {item["id"]: item["name"] for item in template["sheets"]}
    prototypes: dict[str, dict[str, Any]] = {}
    for graph in template["graphs"]:
        sheet_name = sheet_names.get(graph.get("sheet_id"))
        name = graph.get("name", "")
        key: str | None = None
        if sheet_name == "Endurance" and name == "Endurance":
            key = "endurance"
        elif sheet_name == "PUND loops":
            key = "pund_loop"
        elif sheet_name == "PUND metrics" and name == "PUND Ec vs cycle":
            key = "pund_ec"
        elif sheet_name == "PUND metrics" and name == "PUND Pr vs cycle":
            key = "pund_pr"
        elif sheet_name == "PV loops" and name.startswith("PV_IV_"):
            key = "pv_iv"
        elif sheet_name == "PV loops" and name.startswith("PV_"):
            key = "pv_loop"
        elif sheet_name == "PV metrics" and name == "PV Ec vs cycle":
            key = "pv_ec"
        elif sheet_name == "PV metrics" and name == "PV Pr vs cycle":
            key = "pv_pr"
        if key and key not in prototypes:
            prototypes[key] = copy.deepcopy(graph)
    expected = {
        "endurance",
        "pund_loop",
        "pund_ec",
        "pund_pr",
        "pv_loop",
        "pv_iv",
        "pv_ec",
        "pv_pr",
    }
    missing = expected.difference(prototypes)
    if missing:
        raise ValueError(f"Template is missing graph prototypes: {sorted(missing)}")
    return prototypes


def viridis_colors(count: int) -> list[str]:
    if count <= 0:
        return []
    return [
        to_hex(colormaps["viridis"](value))
        for value in np.linspace(0.05, 0.90, count)
    ]


def condition_style(index: int) -> tuple[str, str]:
    line_styles = ("solid", "dashed", "dotted", "dashdot")
    markers = ("o", "s", "^", "D", "v", "P")
    return line_styles[index % len(line_styles)], markers[index % len(markers)]


def condition_color(index: int) -> str:
    colors = ("#009E73", "#56B4E9", "#E69F00", "#CC79A7", "#0072B2", "#666666")
    return colors[index % len(colors)]


def clone_graph(
    prototype: dict[str, Any],
    *,
    name: str,
    sheet_id: str,
    series: list[dict[str, Any]],
    checked_y: list[str],
) -> dict[str, Any]:
    graph = copy.deepcopy(prototype)
    graph["id"] = new_id("gr")
    graph["name"] = name
    graph["sheet_id"] = sheet_id
    graph["series_config"] = series
    graph["checked_y"] = checked_y
    return graph


def endurance_series(
    prototype: dict[str, Any], conditions: list[Condition]
) -> tuple[list[dict[str, Any]], list[str]]:
    base_psw, base_qsw = prototype["series_config"][:2]
    series: list[dict[str, Any]] = []
    checked: list[str] = []
    for index, condition in enumerate(conditions):
        token = condition_token(condition.condition_id)
        color = condition_color(index)
        _, marker = condition_style(index)
        for metric, base in (("Psw", base_psw), ("Qsw", base_qsw)):
            item = copy.deepcopy(base)
            item["x"] = f"cycle_{token}"
            item["y"] = f"{metric}_{token}"
            item["label"] = f"{condition.display_name} {metric}"
            item["color"] = color
            item["marker"] = marker
            item["line_style"] = "solid" if metric == "Psw" else "dashed"
            series.append(item)
            checked.append(item["y"])
    return series, checked


def loop_series(
    prototype: dict[str, Any],
    condition: Condition,
    cycles: list[int | float],
    *,
    prefix_x: str,
    prefix_y: str,
) -> tuple[list[dict[str, Any]], list[str]]:
    base = prototype["series_config"][0]
    colors = viridis_colors(len(cycles))
    token = condition_token(condition.condition_id)
    series: list[dict[str, Any]] = []
    checked: list[str] = []
    for cycle, color in zip(cycles, colors, strict=True):
        suffix = cycle_suffix(cycle)
        item = copy.deepcopy(base)
        item["x"] = f"{prefix_x}_{token}_{suffix}"
        item["y"] = f"{prefix_y}_{token}_{suffix}"
        item["label"] = cycle_display(cycle)
        item["color"] = color
        series.append(item)
        checked.append(item["y"])
    return series, checked


def metric_series(
    prototype: dict[str, Any],
    conditions: list[Condition],
    selected_metrics: tuple[str, str, str],
) -> tuple[list[dict[str, Any]], list[str]]:
    bases = prototype["series_config"][: len(METRICS)]
    if len(bases) != len(METRICS):
        raise ValueError("Metric graph prototype does not contain 12 series")
    series: list[dict[str, Any]] = []
    checked: list[str] = []
    for index, condition in enumerate(conditions):
        token = condition_token(condition.condition_id)
        line_style, marker = condition_style(index)
        for metric, base in zip(METRICS, bases, strict=True):
            item = copy.deepcopy(base)
            item["x"] = f"cycle_{token}"
            item["y"] = metric_identifier(metric, token)
            item["label"] = f"{condition.display_name} {metric}"
            item["line_style"] = line_style
            item["marker"] = marker
            series.append(item)
            if metric in selected_metrics:
                checked.append(item["y"])
    return series, checked


def tree_node(
    node_type: str,
    name: str,
    ref_id: str | None = None,
    children: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    return {
        "id": new_id("nd"),
        "type": node_type,
        "name": name,
        "ref_id": ref_id,
        "expanded": True,
        "children": children or [],
    }


def append_sheet(
    sheets: list[dict[str, Any]],
    parent_children: list[dict[str, Any]],
    name: str,
    data: dict[str, Any],
) -> tuple[str, dict[str, Any]]:
    sheet_id = new_id("sh")
    sheets.append({"id": sheet_id, "name": name, "data": data})
    node = tree_node("sheet", name, sheet_id)
    parent_children.append(node)
    return sheet_id, node


def append_graph(
    graphs: list[dict[str, Any]],
    sheet_node: dict[str, Any],
    graph: dict[str, Any],
) -> str:
    graphs.append(graph)
    node = tree_node("graph", graph["name"], graph["id"])
    sheet_node["children"].append(node)
    return node["id"]


def build_project(
    group: ProcessedGroup,
    template: dict[str, Any],
) -> dict[str, Any]:
    prototypes = extract_graph_prototypes(template)
    sheets: list[dict[str, Any]] = []
    graphs: list[dict[str, Any]] = []
    root = tree_node(
        "folder",
        f"HfO2_{group.spec.thickness_label} {group.spec.file_group}",
    )
    metadata_id = new_id("sh")
    sheets.append(
        {
            "id": metadata_id,
            "name": "Metadata",
            "data": table_sheet_payload(group.metadata),
        }
    )
    metadata_node = tree_node("sheet", "Metadata", metadata_id)
    root["children"].append(metadata_node)
    manifest_id = new_id("sh")
    sheets.append(
        {
            "id": manifest_id,
            "name": "File manifest",
            "data": table_sheet_payload(group.manifest),
        }
    )
    manifest_node = tree_node("sheet", "File manifest", manifest_id)
    root["children"].append(manifest_node)
    active_node_id: str | None = None

    for voltage, tables in group.tables.items():
        voltage_node = tree_node("folder", voltage)
        root["children"].append(voltage_node)

        if not tables["Endurance"].empty:
            payload, conditions = endurance_json_payload(group, voltage)
            sheet_id, sheet_node = append_sheet(
                sheets, voltage_node["children"], "Endurance", payload
            )
            series, checked = endurance_series(prototypes["endurance"], conditions)
            graph = clone_graph(
                prototypes["endurance"],
                name="Endurance",
                sheet_id=sheet_id,
                series=series,
                checked_y=checked,
            )
            node_id = append_graph(graphs, sheet_node, graph)
            active_node_id = active_node_id or node_id

        if not tables["PUND loops"].empty:
            payload, cycle_map = pund_loops_json_payload(group, voltage)
            sheet_id, sheet_node = append_sheet(
                sheets, voltage_node["children"], "PUND loops", payload
            )
            for condition in conditions_for_voltage(group, voltage):
                cycles = cycle_map.get(condition.condition_id)
                if not cycles:
                    continue
                series, checked = loop_series(
                    prototypes["pund_loop"],
                    condition,
                    cycles,
                    prefix_x="Vloop",
                    prefix_y="Ploop",
                )
                graph = clone_graph(
                    prototypes["pund_loop"],
                    name=condition.display_name,
                    sheet_id=sheet_id,
                    series=series,
                    checked_y=checked,
                )
                node_id = append_graph(graphs, sheet_node, graph)
                active_node_id = active_node_id or node_id

        if not tables["PUND metrics"].empty:
            payload, conditions = metrics_json_payload(
                group, voltage, "PUND metrics"
            )
            sheet_id, sheet_node = append_sheet(
                sheets, voltage_node["children"], "PUND metrics", payload
            )
            for key, name, selected in (
                ("pund_ec", "PUND Ec vs cycle", ("Ec+", "Ec-", "2Ec")),
                ("pund_pr", "PUND Pr vs cycle", ("Pr+", "Pr-", "2Pr")),
            ):
                series, checked = metric_series(prototypes[key], conditions, selected)
                graph = clone_graph(
                    prototypes[key],
                    name=name,
                    sheet_id=sheet_id,
                    series=series,
                    checked_y=checked,
                )
                node_id = append_graph(graphs, sheet_node, graph)
                active_node_id = active_node_id or node_id

        if not tables["PV loops"].empty:
            payload, cycle_map = pv_loops_json_payload(group, voltage)
            sheet_id, sheet_node = append_sheet(
                sheets, voltage_node["children"], "PV loops", payload
            )
            for condition in conditions_for_voltage(group, voltage):
                cycles = cycle_map.get(condition.condition_id)
                if not cycles:
                    continue
                for key, graph_prefix, y_prefix in (
                    ("pv_loop", "PV_", "P"),
                    ("pv_iv", "PV_IV_", "J"),
                ):
                    series, checked = loop_series(
                        prototypes[key],
                        condition,
                        cycles,
                        prefix_x="V",
                        prefix_y=y_prefix,
                    )
                    graph = clone_graph(
                        prototypes[key],
                        name=f"{graph_prefix}{condition.display_name}",
                        sheet_id=sheet_id,
                        series=series,
                        checked_y=checked,
                    )
                    node_id = append_graph(graphs, sheet_node, graph)
                    active_node_id = active_node_id or node_id

        if not tables["PV metrics"].empty:
            payload, conditions = metrics_json_payload(group, voltage, "PV metrics")
            sheet_id, sheet_node = append_sheet(
                sheets, voltage_node["children"], "PV metrics", payload
            )
            for key, name, selected in (
                ("pv_ec", "PV Ec vs cycle", ("Ec+", "Ec-", "2Ec")),
                ("pv_pr", "PV Pr vs cycle", ("Pr+", "Pr-", "2Pr")),
            ):
                series, checked = metric_series(prototypes[key], conditions, selected)
                graph = clone_graph(
                    prototypes[key],
                    name=name,
                    sheet_id=sheet_id,
                    series=series,
                    checked_y=checked,
                )
                node_id = append_graph(graphs, sheet_node, graph)
                active_node_id = active_node_id or node_id

    return {
        "schema_version": 3,
        "active_node_id": active_node_id or metadata_node["id"],
        "sheets": sheets,
        "graphs": graphs,
        "tree": root,
    }


def write_project(project: dict[str, Any], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8") as handle:
        json.dump(
            project,
            handle,
            ensure_ascii=False,
            indent=2,
            allow_nan=False,
        )
        handle.write("\n")


def validate_project_payload(project: dict[str, Any]) -> None:
    if project.get("schema_version") != 3:
        raise ValueError("Pubfig schema_version must be 3")
    sheet_ids = [item["id"] for item in project["sheets"]]
    graph_ids = [item["id"] for item in project["graphs"]]
    if len(sheet_ids) != len(set(sheet_ids)):
        raise ValueError("Duplicate sheet IDs")
    if len(graph_ids) != len(set(graph_ids)):
        raise ValueError("Duplicate graph IDs")
    sheet_columns = {
        item["id"]: set(item["data"]["columns"]) for item in project["sheets"]
    }
    for graph in project["graphs"]:
        if graph["sheet_id"] not in sheet_columns:
            raise ValueError(f"Graph {graph['id']} refers to a missing sheet")
        columns = sheet_columns[graph["sheet_id"]]
        series_y = {item["y"] for item in graph["series_config"]}
        for item in graph["series_config"]:
            if item["x"] not in columns or item["y"] not in columns:
                raise ValueError(
                    f"Graph {graph['name']!r} series refers to missing columns: "
                    f"{item['x']!r}, {item['y']!r}"
                )
        if not set(graph["checked_y"]).issubset(series_y):
            raise ValueError(f"Graph {graph['name']!r} has invalid checked_y")

    node_ids: list[str] = []
    node_refs: list[tuple[str, str | None]] = []

    def walk(node: dict[str, Any]) -> None:
        node_ids.append(node["id"])
        node_refs.append((node["type"], node.get("ref_id")))
        for child in node.get("children", []):
            walk(child)

    walk(project["tree"])
    if len(node_ids) != len(set(node_ids)):
        raise ValueError("Duplicate tree node IDs")
    if project["active_node_id"] not in set(node_ids):
        raise ValueError("active_node_id is not a tree node")
    for node_type, ref_id in node_refs:
        if node_type == "sheet" and ref_id not in set(sheet_ids):
            raise ValueError("Sheet tree node has invalid ref_id")
        if node_type == "graph" and ref_id not in set(graph_ids):
            raise ValueError("Graph tree node has invalid ref_id")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source-root",
        type=Path,
        default=Path("/Volumes/ESD-USB/Kiethley/09012026"),
        help="Campaign directory containing HfO2_5nm and HfO2_10nm",
    )
    parser.add_argument(
        "--template-json",
        type=Path,
        default=Path(
            "/Users/ryoo/Desktop/LDRD/Manuscript/HfO2/Json/"
            "0826_5nm_duration/HfO2_5nm_5.0ms.json"
        ),
    )
    parser.add_argument(
        "--data-output-root",
        type=Path,
        required=True,
        help="Root that will receive HfO2_5nm/ and HfO2_10nm/",
    )
    parser.add_argument(
        "--json-output-root",
        type=Path,
        required=True,
        help="Root that will receive 5nm/ and 10nm/",
    )
    parser.add_argument("--date-token", default=DATE_TOKEN)
    parser.add_argument("--date-display", default=DATE_DISPLAY)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    with args.template_json.open(encoding="utf-8") as handle:
        template = json.load(handle)
    summaries: list[dict[str, Any]] = []
    for spec in DEFAULT_GROUPS:
        group = process_group(
            args.source_root,
            spec,
            date_token=args.date_token,
            date_display=args.date_display,
            script_path=Path(__file__),
        )
        workbook_path = (
            args.data_output_root
            / f"HfO2_{spec.thickness_label}"
            / f"{spec.basename}.xlsx"
        )
        project_path = (
            args.json_output_root
            / spec.json_subdir
            / f"{spec.basename}.json"
        )
        write_workbook(group, workbook_path)
        project = build_project(group, template)
        validate_project_payload(project)
        write_project(project, project_path)
        summaries.append(
            {
                "group": spec.basename,
                "source_files": len(group.manifest),
                "conditions": [item.condition_id for item in group.conditions],
                "warnings": int((group.manifest["status"] == "warning").sum()),
                "errors": int((group.manifest["status"] == "error").sum()),
                "workbook": str(workbook_path),
                "project": str(project_path),
                "sheets": len(project["sheets"]),
                "graphs": len(project["graphs"]),
            }
        )
    print(json.dumps(summaries, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
