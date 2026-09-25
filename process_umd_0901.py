#!/usr/bin/env python3
"""Process the 2026-09-01 UMD Keithley campaign into XLSX and Pubfig JSON.

The UMD source tree is intentionally handled independently of the HfO2 batch
layout: modes are detected from the Data-sheet columns, bare ``endurance.xls``
files are supported, ``col37/4V.xls`` is retained as a named pretest, and the
misspelled source directory ``Archieve`` is exported as one voltage-series
dataset rather than being misread as fatigue cycles.

Polarization and current density use the 400 um2 electrode area from the
user-supplied 2026-08-26 reference.  The UMD thickness is not present in the
raw workbooks, so field values are deliberately left blank instead of
silently inheriting the reference sample's 5 nm thickness.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import math
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

import process_hfo2_0901 as common


DATE_TOKEN = "0901"
DATE_DISPLAY = "09/01"
AREA_UM2 = 400.0
MODE_ORDER = {"Endurance": 0, "PUND": 1, "PV": 2}
MODE_SIGNATURES = {
    "Endurance": {"pundEndurance", "iteration", "Psw", "Qsw"},
    "PUND": {"pundTest", "V", "I", "t"},
    "PV": {"doubleSweepSeg", "Vforce", "Imeas", "Charge"},
}

EXTRA_COLUMNS = (
    "measurement_id",
    "measurement_label",
    "coordinate_kind",
    "coordinate_value",
    "drive_voltage_V",
    "executed_at",
)
ENDURANCE_COLUMNS = (
    "condition",
    "file",
    *EXTRA_COLUMNS,
    *common.ENDURANCE_COLUMNS[2:],
)
PUND_LOOP_COLUMNS = (
    "condition",
    "cycle",
    "file",
    *EXTRA_COLUMNS,
    *common.PUND_LOOP_COLUMNS[3:],
)
PV_LOOP_COLUMNS = (
    "condition",
    "cycle",
    "file",
    *EXTRA_COLUMNS,
    *common.PV_LOOP_COLUMNS[3:],
)
METRIC_COLUMNS = (
    "condition",
    "cycle",
    "file",
    *EXTRA_COLUMNS,
    *common.METRICS,
)
MANIFEST_COLUMNS = (
    *common.MANIFEST_COLUMNS,
    "dataset",
    "sample_id",
    "column_or_group",
    "coordinate_kind",
    "coordinate_value",
    "coordinate_source",
    "drive_voltage_V",
    "executed_at",
    "mode_source",
    "sha256",
)


@dataclass(frozen=True)
class DatasetSpec:
    sample_id: str
    group_name: str
    relative_dir: str
    kind: str = "cycle"

    @property
    def basename(self) -> str:
        return f"{self.sample_id}_{self.group_name}"

    @property
    def display_name(self) -> str:
        if self.kind == "voltage_sweep":
            return f"{self.sample_id} Archieve voltage series ({DATE_DISPLAY})"
        return f"{self.sample_id} {self.group_name} ({DATE_DISPLAY})"


@dataclass(frozen=True)
class Measurement:
    path: Path
    mode: str
    measurement_id: str
    label: str
    coordinate_kind: str
    coordinate_value: float | int | None
    cycle: float | int | None
    coordinate_source: str
    drive_voltage_v: float | None
    executed_at: str | None
    max_loops: int | None


@dataclass
class DatasetResult:
    spec: DatasetSpec
    source_dir: Path
    measurements: list[Measurement]
    metadata: pd.DataFrame
    manifest: pd.DataFrame
    tables: dict[str, pd.DataFrame]


def numeric_value(value: Any) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def normalized_number(value: float | int | None) -> float | int | None:
    if value is None:
        return None
    numeric = float(value)
    return int(numeric) if numeric.is_integer() else numeric


def safe_token(value: str) -> str:
    token = re.sub(r"[^A-Za-z0-9_-]+", "_", value).strip("_")
    return token or "measurement"


def hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def setting_value(settings: pd.DataFrame, key: str, column: int) -> Any:
    if settings.empty or settings.shape[1] <= column:
        return None
    keys = settings.iloc[:, 0].astype(str).str.strip()
    rows = settings[keys == key]
    if rows.empty:
        return None
    value = rows.iloc[0, column]
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    return value


def classify_mode(path: Path) -> str:
    frame = common.read_excel_quiet(path, "Data")
    columns = {str(item).strip() for item in frame.columns}
    matches = [
        mode for mode, signature in MODE_SIGNATURES.items() if signature <= columns
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one mode signature for {path}, found {matches or 'none'}"
        )
    return matches[0]


def parse_filename_cycle(path: Path) -> int | float | None:
    if not re.fullmatch(r"(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", path.stem):
        return None
    return normalized_number(float(path.stem))


def discover_dataset_specs(source_root: Path) -> list[DatasetSpec]:
    specs: list[DatasetSpec] = []
    for sample_id in ("UMD_13", "UMD_14"):
        sample_dir = source_root / sample_id
        if not sample_dir.is_dir():
            raise FileNotFoundError(f"Missing sample directory: {sample_dir}")
        columns = [
            item
            for item in sample_dir.iterdir()
            if item.is_dir() and re.fullmatch(r"col\d+(?:-\d+)?", item.name, re.I)
        ]
        columns.sort(
            key=lambda item: (
                int(re.search(r"\d+", item.name).group()),  # type: ignore[union-attr]
                item.name,
            )
        )
        for column in columns:
            specs.append(DatasetSpec(sample_id, column.name, f"{sample_id}/{column.name}"))

        archive_dirs = [
            item
            for item in sample_dir.iterdir()
            if item.is_dir() and item.name.lower() in {"archive", "archieve"}
        ]
        if len(archive_dirs) > 1:
            raise ValueError(f"Multiple archive directories found in {sample_dir}")
        if archive_dirs:
            archive = archive_dirs[0]
            specs.append(
                DatasetSpec(
                    sample_id,
                    archive.name,
                    f"{sample_id}/{archive.name}",
                    kind="voltage_sweep",
                )
            )
    return specs


def discover_measurements(source_root: Path, spec: DatasetSpec) -> list[Measurement]:
    source_dir = source_root / spec.relative_dir
    paths = sorted(source_dir.rglob("*.xls"))
    if not paths:
        raise ValueError(f"No XLS files found in {source_dir}")

    measurements: list[Measurement] = []
    seen: set[tuple[str, str]] = set()
    for path in paths:
        mode = classify_mode(path)
        settings = common.read_excel_quiet(path, "Settings", header=None)
        executed = setting_value(settings, "Last Executed", 1)
        executed_at = str(executed).strip() if executed is not None else None
        drive_key = "V1" if mode == "PV" else "Vp"
        drive_voltage = numeric_value(setting_value(settings, drive_key, 3))
        max_loops_value = numeric_value(setting_value(settings, "max_loops", 3))
        max_loops = int(max_loops_value) if max_loops_value is not None else None
        file_cycle = parse_filename_cycle(path)

        if spec.kind == "voltage_sweep":
            if mode != "PV":
                raise ValueError(f"Archive contains non-PV file: {path}")
            if drive_voltage is None:
                raise ValueError(f"Archive PV file has no Settings.V1: {path}")
            coordinate = normalized_number(abs(drive_voltage))
            voltage_text = f"{float(coordinate):g}V"
            measurement_id = f"drive_{safe_token(voltage_text)}"
            label = f"{float(coordinate):g} V sweep"
            coordinate_kind = "drive_voltage"
            coordinate_source = "Settings.V1"
            cycle = None
        elif file_cycle is not None:
            coordinate = file_cycle
            measurement_id = f"cycle_{safe_token(common.cycle_suffix(file_cycle))}"
            label = common.cycle_display(file_cycle)
            coordinate_kind = "cycle"
            coordinate_source = "numeric filename stem"
            cycle = file_cycle
        elif mode == "Endurance" and max_loops is not None:
            coordinate = max_loops
            measurement_id = "endurance_full_run"
            label = f"Endurance through {max_loops:g} cycles"
            coordinate_kind = "endurance_block"
            coordinate_source = "Settings.max_loops"
            cycle = None
        elif mode == "PV" and re.fullmatch(r"\d+(?:\.\d+)?V", path.stem, re.I):
            coordinate = None
            measurement_id = f"named_{safe_token(path.stem)}"
            voltage_text = path.stem[:-1]
            label = f"{voltage_text} V pretest"
            coordinate_kind = "named_pretest"
            coordinate_source = "non-cycle filename and execution order"
            cycle = None
        else:
            raise ValueError(f"Cannot assign a measurement coordinate to {path}")

        unique_key = (mode, measurement_id)
        if unique_key in seen:
            raise ValueError(f"Duplicate {mode} measurement ID {measurement_id} in {source_dir}")
        seen.add(unique_key)
        measurements.append(
            Measurement(
                path=path,
                mode=mode,
                measurement_id=measurement_id,
                label=label,
                coordinate_kind=coordinate_kind,
                coordinate_value=coordinate,
                cycle=cycle,
                coordinate_source=coordinate_source,
                drive_voltage_v=drive_voltage,
                executed_at=executed_at,
                max_loops=max_loops,
            )
        )

    measurements.sort(key=measurement_sort_key)
    return measurements


def executed_sort_value(value: str | None) -> str:
    if not value:
        return "9999"
    for pattern in ("%m/%d/%Y %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(value, pattern).isoformat()
        except ValueError:
            continue
    return value


def measurement_sort_key(item: Measurement) -> tuple[Any, ...]:
    if item.coordinate_kind == "named_pretest":
        coordinate_order = 0
        coordinate = -math.inf
    else:
        coordinate_order = 1
        coordinate = (
            float(item.coordinate_value)
            if item.coordinate_value is not None
            else math.inf
        )
    return (
        MODE_ORDER[item.mode],
        coordinate_order,
        coordinate,
        executed_sort_value(item.executed_at),
        item.path.name,
    )


def empty_frame(columns: Iterable[str]) -> pd.DataFrame:
    return pd.DataFrame(columns=list(columns))


def add_measurement_columns(frame: pd.DataFrame, item: Measurement) -> pd.DataFrame:
    frame["measurement_id"] = item.measurement_id
    frame["measurement_label"] = item.label
    frame["coordinate_kind"] = item.coordinate_kind
    frame["coordinate_value"] = item.coordinate_value
    frame["drive_voltage_V"] = item.drive_voltage_v
    frame["executed_at"] = item.executed_at
    return frame


def metric_record(
    raw_metric: dict[str, Any], spec: DatasetSpec, item: Measurement
) -> dict[str, Any]:
    result: dict[str, Any] = {
        "condition": spec.basename,
        "cycle": item.cycle,
        "file": item.path.name,
        "measurement_id": item.measurement_id,
        "measurement_label": item.label,
        "coordinate_kind": item.coordinate_kind,
        "coordinate_value": item.coordinate_value,
        "drive_voltage_V": item.drive_voltage_v,
        "executed_at": item.executed_at,
    }
    result.update({metric: raw_metric.get(metric) for metric in common.METRICS})
    return result


def source_relative_path(path: Path, source_root: Path) -> str:
    try:
        return path.relative_to(source_root.parent).as_posix()
    except ValueError:
        return path.as_posix()


def manifest_record(
    source_root: Path,
    spec: DatasetSpec,
    item: Measurement,
    *,
    input_rows: int,
    output_rows: int,
    status: str,
    warning: str | None,
) -> dict[str, Any]:
    pulse_condition = (
        "voltage_sweep"
        if spec.kind == "voltage_sweep"
        else (
            f"{float(item.drive_voltage_v):g}V"
            if item.drive_voltage_v is not None
            else None
        )
    )
    return {
        "pulse_condition": pulse_condition,
        "condition": spec.basename,
        "mode": item.mode,
        "cycle_file_value": item.cycle,
        "source_file": item.path.name,
        "source_relative_path": source_relative_path(item.path, source_root),
        "input_rows": input_rows,
        "output_rows": output_rows,
        "status": status,
        "warning": warning,
        "size_bytes": item.path.stat().st_size,
        "dataset": spec.basename,
        "sample_id": spec.sample_id,
        "column_or_group": spec.group_name,
        "coordinate_kind": item.coordinate_kind,
        "coordinate_value": item.coordinate_value,
        "coordinate_source": item.coordinate_source,
        "drive_voltage_V": item.drive_voltage_v,
        "executed_at": item.executed_at,
        "mode_source": "Data-sheet column signature",
        "sha256": hash_file(item.path),
    }


def process_dataset(
    source_root: Path,
    spec: DatasetSpec,
    script_path: Path,
) -> DatasetResult:
    source_dir = source_root / spec.relative_dir
    measurements = discover_measurements(source_root, spec)
    tables = {
        "Endurance": empty_frame(ENDURANCE_COLUMNS),
        "PUND loops": empty_frame(PUND_LOOP_COLUMNS),
        "PUND metrics": empty_frame(METRIC_COLUMNS),
        "PV loops": empty_frame(PV_LOOP_COLUMNS),
        "PV metrics": empty_frame(METRIC_COLUMNS),
    }
    blocks: dict[str, list[pd.DataFrame]] = {name: [] for name in tables}
    manifest_rows: list[dict[str, Any]] = []

    previous_end = 0
    for item in measurements:
        input_rows = 0
        output_rows = 0
        status = "ok"
        warning: str | None = None
        try:
            if item.mode == "Endurance":
                data = common.read_excel_quiet(item.path, "Data")
                input_rows = len(data)
                settings = common.read_excel_quiet(item.path, "Settings", header=None)
                max_loops = int(common.get_setting(settings, "max_loops"))
                missing = {"iteration", "Psw", "Qsw"}.difference(data.columns)
                if missing:
                    raise ValueError(f"Missing Endurance columns: {sorted(missing)}")
                block = data.copy()
                block.insert(0, "file", item.path.name)
                iteration = pd.to_numeric(block["iteration"], errors="coerce")
                block.insert(1, "global_cycle", previous_end + iteration)
                block["plot_cycle"] = block["global_cycle"].clip(lower=1)
                block["axis_segment"] = 0
                block["axis_cycle"] = block["plot_cycle"]
                block["Psw_uC_cm2"] = common.charge_to_polarization(
                    pd.to_numeric(block["Psw"], errors="coerce"), AREA_UM2
                )
                block["Qsw_uC_cm2"] = common.charge_to_polarization(
                    pd.to_numeric(block["Qsw"], errors="coerce"), AREA_UM2
                )
                block.insert(0, "condition", spec.basename)
                block = add_measurement_columns(block, item)
                block = block.reindex(columns=ENDURANCE_COLUMNS)
                output_rows = len(block)
                warning = common.endurance_collapse_warning(block)
                if warning:
                    status = "warning"
                blocks["Endurance"].append(block)
                previous_end += max_loops

            elif item.mode == "PUND":
                loop, input_rows = common.calculate_pund_file(item.path)
                output_rows = len(loop)
                metric = common.pund_metric_row(
                    loop,
                    item.cycle if item.cycle is not None else 0,
                    item.path.name,
                )
                loop.insert(0, "file", item.path.name)
                loop.insert(0, "cycle", item.cycle)
                loop.insert(0, "condition", spec.basename)
                loop = add_measurement_columns(loop, item)
                blocks["PUND loops"].append(loop.reindex(columns=PUND_LOOP_COLUMNS))
                blocks["PUND metrics"].append(
                    pd.DataFrame([metric_record(metric, spec, item)]).reindex(
                        columns=METRIC_COLUMNS
                    )
                )

            elif item.mode == "PV":
                metric_coordinate = (
                    item.coordinate_value if item.coordinate_value is not None else 0
                )
                loop, metric, input_rows = common.calculate_pv_file(
                    item.path,
                    metric_coordinate,
                    float("nan"),
                )
                output_rows = len(loop)
                loop.insert(0, "file", item.path.name)
                loop.insert(0, "cycle", item.cycle)
                loop.insert(0, "condition", spec.basename)
                loop = add_measurement_columns(loop, item)
                blocks["PV loops"].append(loop.reindex(columns=PV_LOOP_COLUMNS))
                blocks["PV metrics"].append(
                    pd.DataFrame([metric_record(metric, spec, item)]).reindex(
                        columns=METRIC_COLUMNS
                    )
                )
                if item.coordinate_kind == "named_pretest":
                    warning = (
                        "non-cycle PV filename retained as an earlier named pretest; "
                        "loop and metrics are preserved, but it is omitted from the "
                        "cycle-metric graph"
                    )
                    status = "warning"
            else:  # pragma: no cover - discovery prevents this
                raise ValueError(f"Unsupported mode: {item.mode}")
        except Exception as exc:
            status = "error"
            warning = f"{type(exc).__name__}: {exc}"

        manifest_rows.append(
            manifest_record(
                source_root,
                spec,
                item,
                input_rows=input_rows,
                output_rows=output_rows,
                status=status,
                warning=warning,
            )
        )

    for name, frames in blocks.items():
        if frames:
            tables[name] = pd.concat(frames, ignore_index=True)

    manifest = pd.DataFrame(manifest_rows).reindex(columns=MANIFEST_COLUMNS)
    manifest["_mode_order"] = manifest["mode"].map(MODE_ORDER)
    manifest = (
        manifest.sort_values(
            ["_mode_order", "coordinate_value", "executed_at", "source_file"],
            kind="stable",
            na_position="first",
        )
        .drop(columns="_mode_order")
        .reset_index(drop=True)
    )
    metadata = build_metadata(spec, source_dir, measurements, manifest, script_path)
    return DatasetResult(spec, source_dir, measurements, metadata, manifest, tables)


def build_metadata(
    spec: DatasetSpec,
    source_dir: Path,
    measurements: list[Measurement],
    manifest: pd.DataFrame,
    script_path: Path,
) -> pd.DataFrame:
    modes = [
        mode for mode in MODE_ORDER if any(item.mode == mode for item in measurements)
    ]
    missing_modes = [mode for mode in MODE_ORDER if mode not in modes]
    drive_values = sorted(
        {
            normalized_number(abs(item.drive_voltage_v))
            for item in measurements
            if item.drive_voltage_v is not None
        },
        key=float,
    )
    warning_count = int((manifest["status"] == "warning").sum())
    error_count = int((manifest["status"] == "error").sum())
    named = [item.label for item in measurements if item.coordinate_kind == "named_pretest"]
    coordinate_description = (
        "Drive voltage from Settings.V1; files are independent amplitude sweeps, not fatigue cycles."
        if spec.kind == "voltage_sweep"
        else "Fatigue cycle from numeric filenames; named pretests remain explicitly unindexed."
    )
    rows: list[tuple[str, Any]] = [
        ("dataset", spec.basename),
        ("sample_id", spec.sample_id),
        ("column_or_group", spec.group_name),
        ("measurement_date", "2026-09-01"),
        ("material", "Not encoded in the raw workbook; sample ID preserved"),
        ("nominal_thickness_nm", None),
        (
            "thickness_source",
            "Not available in source path or XLS Settings; 5 nm reference thickness was not inherited.",
        ),
        (
            "field_conversion",
            "Not calculated; Field_MV_cm is blank because thickness is unknown.",
        ),
        ("source_root", str(source_dir)),
        ("source_file_count", len(manifest)),
        ("source_modes", ", ".join(modes)),
        ("missing_modes", ", ".join(missing_modes) if missing_modes else None),
        ("drive_voltage_values_V", ", ".join(f"{float(v):g}" for v in drive_values)),
        ("measurement_coordinate", coordinate_description),
        ("electrode_area_um2", AREA_UM2),
        (
            "electrode_area_source",
            "User-supplied 2026-08-26 processing reference / analyzer GUI default; raw XLS has no area field.",
        ),
        ("metric_ec_axis", "Voltage (V), matching the supplied reference"),
        (
            "pund_curve_definition",
            "One loop per cycle: P-U (2000 points) followed by N-D (2000 points) in Pubfig JSON.",
        ),
        ("named_measurements", "; ".join(named) if named else None),
        ("warning_file_count", warning_count),
        ("error_file_count", error_count),
        (
            "organization_note",
            "Each column is exported as one independent device; Archieve is one multi-voltage dataset.",
        ),
        (
            "source_copy_note",
            "Desktop organized source selected after SHA-256 comparison showed all 49 UMD XLS files byte-identical to the USB source.",
        ),
        ("generated_utc", datetime.now(timezone.utc).isoformat()),
        ("analysis_script", str(script_path.resolve())),
        (
            "analysis_reference",
            str((script_path.parent / "ferro_cycle_analyzer_fixed.py").resolve()),
        ),
        (
            "format_template",
            "/Users/ryoo/Desktop/LDRD/Manuscript/HfO2/Json/0826_5nm_duration/HfO2_5nm_5.0ms.json",
        ),
    ]
    return pd.DataFrame(rows, columns=["key", "value"])


def write_workbook(result: DatasetResult, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        result.metadata.to_excel(writer, sheet_name="Metadata", index=False)
        result.manifest.to_excel(writer, sheet_name="File_manifest", index=False)
        for name in (
            "Endurance",
            "PUND loops",
            "PUND metrics",
            "PV loops",
            "PV metrics",
        ):
            frame = result.tables[name]
            if not frame.empty:
                frame.to_excel(writer, sheet_name=name.replace(" ", "_"), index=False)

        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for column_cells in worksheet.columns:
                column_letter = column_cells[0].column_letter
                maximum = max(
                    (len(str(cell.value)) for cell in column_cells if cell.value is not None),
                    default=0,
                )
                worksheet.column_dimensions[column_letter].width = min(
                    max(maximum + 2, 12), 55
                )


def measurement_items(result: DatasetResult, mode: str) -> list[Measurement]:
    return [item for item in result.measurements if item.mode == mode]


def endurance_payload(result: DatasetResult) -> dict[str, Any]:
    frame = result.tables["Endurance"]
    token = safe_token(result.spec.basename)
    return common.wide_payload(
        [f"cycle_{token}", f"Psw_{token}", f"Qsw_{token}"],
        ["X", "Y", "Y"],
        [
            f"{result.spec.display_name} cycle",
            f"{result.spec.display_name} Psw",
            f"{result.spec.display_name} Qsw",
        ],
        [
            frame["plot_cycle"].tolist(),
            frame["Psw_uC_cm2"].tolist(),
            frame["Qsw_uC_cm2"].tolist(),
        ],
    )


def loop_payload(
    result: DatasetResult, mode: str
) -> tuple[dict[str, Any], list[tuple[Measurement, str, str, str | None]]]:
    table_name = "PUND loops" if mode == "PUND" else "PV loops"
    frame = result.tables[table_name]
    token = safe_token(result.spec.basename)
    columns: list[str] = []
    roles: list[str] = []
    labels: list[str] = []
    values: list[list[Any]] = []
    bindings: list[tuple[Measurement, str, str, str | None]] = []
    for item in measurement_items(result, mode):
        subset = frame[frame["measurement_id"] == item.measurement_id]
        if subset.empty:
            continue
        suffix = safe_token(item.measurement_id)
        if mode == "PUND":
            x_id = f"Vloop_{token}_{suffix}"
            p_id = f"Ploop_{token}_{suffix}"
            columns.extend([x_id, p_id])
            roles.extend(["X", "Y"])
            labels.extend([item.label, item.label])
            values.extend(
                [
                    [
                        *subset["Voltage_Pos_V"].tolist(),
                        *subset["Voltage_Neg_V"].tolist(),
                    ],
                    [
                        *subset["P_Pos_uC_cm2"].tolist(),
                        *subset["P_Neg_uC_cm2"].tolist(),
                    ],
                ]
            )
            bindings.append((item, x_id, p_id, None))
        else:
            x_id = f"V_{token}_{suffix}"
            p_id = f"P_{token}_{suffix}"
            j_id = f"J_{token}_{suffix}"
            columns.extend([x_id, p_id, j_id])
            roles.extend(["X", "Y", "Y"])
            labels.extend([item.label, item.label, item.label])
            values.extend(
                [
                    subset["Voltage_V"].tolist(),
                    subset["P_uC_cm2"].tolist(),
                    subset["J_A_cm2"].tolist(),
                ]
            )
            bindings.append((item, x_id, p_id, j_id))
    return common.wide_payload(columns, roles, labels, values), bindings


def metrics_payload(
    result: DatasetResult, mode: str
) -> tuple[dict[str, Any], str | None, dict[str, str]]:
    frame = result.tables[f"{mode} metrics"]
    token = safe_token(result.spec.basename)
    indexed = frame[pd.to_numeric(frame["coordinate_value"], errors="coerce").notna()]
    indexed = indexed.copy()
    indexed["coordinate_value"] = pd.to_numeric(
        indexed["coordinate_value"], errors="coerce"
    )
    indexed = indexed.sort_values("coordinate_value", kind="stable")
    columns: list[str] = []
    roles: list[str] = []
    labels: list[str] = []
    values: list[list[Any]] = []
    metric_ids: dict[str, str] = {}
    x_id: str | None = None
    coordinate_label = (
        "Drive voltage (V)" if result.spec.kind == "voltage_sweep" else "Cycle"
    )
    if not indexed.empty:
        x_id = f"coordinate_{token}"
        columns.append(x_id)
        roles.append("X")
        labels.append(coordinate_label)
        values.append(indexed["coordinate_value"].tolist())
        for metric in common.METRICS:
            metric_id = common.metric_identifier(metric, token)
            metric_ids[metric] = metric_id
            columns.append(metric_id)
            roles.append("Y")
            labels.append(f"{result.spec.display_name} {metric}")
            values.append(indexed[metric].tolist())

    unindexed = frame[pd.to_numeric(frame["coordinate_value"], errors="coerce").isna()]
    for _, row in unindexed.iterrows():
        suffix = safe_token(str(row["measurement_id"]))
        label_id = f"unindexed_label_{token}_{suffix}"
        columns.append(label_id)
        roles.append("")
        labels.append("Unindexed measurement")
        values.append([row["measurement_label"]])
        for metric_index, metric in enumerate(common.METRICS):
            # Use the stable metric position rather than a punctuation-stripped
            # token: Ec+ and |Ec+| (likewise Pr variants) must never collapse to
            # the same internal Pubfig column identifier.
            helper_id = f"unindexed_metric_{metric_index:02d}_{token}_{suffix}"
            columns.append(helper_id)
            roles.append("")
            labels.append(f"{row['measurement_label']} {metric}")
            values.append([row[metric]])
    return common.wide_payload(columns, roles, labels, values), x_id, metric_ids


def clone_loop_graph(
    prototype: dict[str, Any],
    *,
    name: str,
    sheet_id: str,
    bindings: list[tuple[Measurement, str, str]],
) -> dict[str, Any]:
    base_series = prototype["series_config"][0]
    colors = common.viridis_colors(len(bindings))
    series: list[dict[str, Any]] = []
    checked: list[str] = []
    for (item, x_id, y_id), color in zip(bindings, colors, strict=True):
        config = copy.deepcopy(base_series)
        config["x"] = x_id
        config["y"] = y_id
        config["label"] = item.label
        config["color"] = color
        series.append(config)
        checked.append(y_id)
    return common.clone_graph(
        prototype,
        name=name,
        sheet_id=sheet_id,
        series=series,
        checked_y=checked,
    )


def clone_metric_graph(
    prototype: dict[str, Any],
    *,
    name: str,
    sheet_id: str,
    x_id: str,
    metric_ids: dict[str, str],
    selected: tuple[str, str, str],
    coordinate_kind: str,
) -> dict[str, Any]:
    bases = prototype["series_config"][: len(common.METRICS)]
    if len(bases) != len(common.METRICS):
        raise ValueError("Metric prototype does not contain all 12 metric styles")
    series: list[dict[str, Any]] = []
    checked: list[str] = []
    for metric, base_series in zip(common.METRICS, bases, strict=True):
        config = copy.deepcopy(base_series)
        config["x"] = x_id
        config["y"] = metric_ids[metric]
        config["label"] = metric
        series.append(config)
        if metric in selected:
            checked.append(config["y"])
    graph = common.clone_graph(
        prototype,
        name=name,
        sheet_id=sheet_id,
        series=series,
        checked_y=checked,
    )
    if coordinate_kind == "drive_voltage":
        graph["plot_config"]["x_label"] = "Drive voltage (V)"
        graph["plot_config"]["x_scale"] = "linear"
        graph["plot_config"]["x_min"] = None
        graph["plot_config"]["x_max"] = None
    return graph


def build_project(result: DatasetResult, template: dict[str, Any]) -> dict[str, Any]:
    prototypes = common.extract_graph_prototypes(template)
    sheets: list[dict[str, Any]] = []
    graphs: list[dict[str, Any]] = []
    root = common.tree_node("folder", result.spec.display_name)

    metadata_id, metadata_node = common.append_sheet(
        sheets,
        root["children"],
        "Metadata",
        common.table_sheet_payload(result.metadata),
    )
    del metadata_id
    common.append_sheet(
        sheets,
        root["children"],
        "File manifest",
        common.table_sheet_payload(result.manifest),
    )
    active_node_id: str | None = None

    if not result.tables["Endurance"].empty:
        sheet_id, sheet_node = common.append_sheet(
            sheets, root["children"], "Endurance", endurance_payload(result)
        )
        token = safe_token(result.spec.basename)
        condition = common.Condition(
            condition_id=result.spec.basename,
            display_name=result.spec.display_name,
            voltage="4V",
            voltage_value=4.0,
            column_name=result.spec.group_name,
            leaf_relative=result.spec.relative_dir,
        )
        series, checked = common.endurance_series(prototypes["endurance"], [condition])
        graph = common.clone_graph(
            prototypes["endurance"],
            name="Endurance",
            sheet_id=sheet_id,
            series=series,
            checked_y=checked,
        )
        # endurance_series uses the same sanitized condition token as our payload.
        assert all(config["x"] == f"cycle_{token}" for config in series)
        active_node_id = common.append_graph(graphs, sheet_node, graph)

    if not result.tables["PUND loops"].empty:
        payload, bindings = loop_payload(result, "PUND")
        sheet_id, sheet_node = common.append_sheet(
            sheets, root["children"], "PUND loops", payload
        )
        graph = clone_loop_graph(
            prototypes["pund_loop"],
            name=f"PUND_{result.spec.display_name}",
            sheet_id=sheet_id,
            bindings=[(item, x_id, p_id) for item, x_id, p_id, _ in bindings],
        )
        node_id = common.append_graph(graphs, sheet_node, graph)
        active_node_id = active_node_id or node_id

    if not result.tables["PUND metrics"].empty:
        payload, x_id, metric_ids = metrics_payload(result, "PUND")
        sheet_id, sheet_node = common.append_sheet(
            sheets, root["children"], "PUND metrics", payload
        )
        if x_id is not None:
            for key, name, selected in (
                ("pund_ec", "PUND Ec vs cycle", ("Ec+", "Ec-", "2Ec")),
                ("pund_pr", "PUND Pr vs cycle", ("Pr+", "Pr-", "2Pr")),
            ):
                graph = clone_metric_graph(
                    prototypes[key],
                    name=name,
                    sheet_id=sheet_id,
                    x_id=x_id,
                    metric_ids=metric_ids,
                    selected=selected,
                    coordinate_kind="cycle",
                )
                node_id = common.append_graph(graphs, sheet_node, graph)
                active_node_id = active_node_id or node_id

    if not result.tables["PV loops"].empty:
        payload, bindings = loop_payload(result, "PV")
        sheet_id, sheet_node = common.append_sheet(
            sheets, root["children"], "PV loops", payload
        )
        for key, prefix, binding_index in (
            ("pv_loop", "PV_", 2),
            ("pv_iv", "PV_IV_", 3),
        ):
            graph_bindings = [
                (item, x_id, binding[binding_index])
                for binding in bindings
                for item, x_id in [(binding[0], binding[1])]
                if binding[binding_index] is not None
            ]
            graph = clone_loop_graph(
                prototypes[key],
                name=f"{prefix}{result.spec.display_name}",
                sheet_id=sheet_id,
                bindings=graph_bindings,
            )
            node_id = common.append_graph(graphs, sheet_node, graph)
            active_node_id = active_node_id or node_id

    if not result.tables["PV metrics"].empty:
        payload, x_id, metric_ids = metrics_payload(result, "PV")
        sheet_id, sheet_node = common.append_sheet(
            sheets, root["children"], "PV metrics", payload
        )
        if x_id is not None:
            coordinate_kind = (
                "drive_voltage" if result.spec.kind == "voltage_sweep" else "cycle"
            )
            suffix = "drive voltage" if coordinate_kind == "drive_voltage" else "cycle"
            for key, metric_name, selected in (
                ("pv_ec", "Ec", ("Ec+", "Ec-", "2Ec")),
                ("pv_pr", "Pr", ("Pr+", "Pr-", "2Pr")),
            ):
                graph = clone_metric_graph(
                    prototypes[key],
                    name=f"PV {metric_name} vs {suffix}",
                    sheet_id=sheet_id,
                    x_id=x_id,
                    metric_ids=metric_ids,
                    selected=selected,
                    coordinate_kind=coordinate_kind,
                )
                node_id = common.append_graph(graphs, sheet_node, graph)
                active_node_id = active_node_id or node_id

    return {
        "schema_version": 3,
        "active_node_id": active_node_id or metadata_node["id"],
        "sheets": sheets,
        "graphs": graphs,
        "tree": root,
    }


def validate_result(result: DatasetResult) -> None:
    if result.manifest.empty:
        raise ValueError(f"Empty manifest for {result.spec.basename}")
    error_rows = result.manifest[result.manifest["status"] == "error"]
    if not error_rows.empty:
        details = error_rows[["source_relative_path", "warning"]].to_dict("records")
        raise ValueError(f"Processing errors for {result.spec.basename}: {details}")
    expected = {item.path.resolve() for item in result.measurements}
    if len(result.manifest) != len(expected):
        raise ValueError(
            f"Manifest count mismatch for {result.spec.basename}: "
            f"{len(result.manifest)} != {len(expected)}"
        )
    for mode, table_name in (
        ("Endurance", "Endurance"),
        ("PUND", "PUND loops"),
        ("PV", "PV loops"),
    ):
        successful = result.manifest[
            (result.manifest["mode"] == mode)
            & (result.manifest["status"].isin(["ok", "warning"]))
        ]
        expected_rows = int(successful["output_rows"].sum())
        actual_rows = len(result.tables[table_name])
        if expected_rows != actual_rows:
            raise ValueError(
                f"{result.spec.basename} {mode} row mismatch: "
                f"{expected_rows} != {actual_rows}"
            )


def validate_project(project: dict[str, Any]) -> None:
    common.validate_project_payload(project)
    for sheet in project["sheets"]:
        columns = sheet["data"]["columns"]
        if len(columns) != len(set(columns)):
            duplicates = sorted(
                {column for column in columns if columns.count(column) > 1}
            )
            raise ValueError(
                f"Duplicate Pubfig column IDs in sheet {sheet['name']!r}: "
                f"{duplicates}"
            )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source-root",
        type=Path,
        default=Path(
            "/Users/ryoo/Desktop/LDRD/Electrical data/Kiethley/09012026"
        ),
    )
    parser.add_argument(
        "--template-json",
        type=Path,
        default=Path(
            "/Users/ryoo/Desktop/LDRD/Manuscript/HfO2/Json/"
            "0826_5nm_duration/HfO2_5nm_5.0ms.json"
        ),
    )
    parser.add_argument(
        "--output-root",
        type=Path,
        default=Path(
            "/Users/ryoo/Desktop/LDRD/Electrical data/Kiethley/"
            "09012026/UMD_processed"
        ),
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    with args.template_json.open(encoding="utf-8") as handle:
        template = json.load(handle)

    summaries: list[dict[str, Any]] = []
    source_paths: set[Path] = set()
    for spec in discover_dataset_specs(args.source_root):
        result = process_dataset(args.source_root, spec, Path(__file__))
        validate_result(result)
        current_paths = {item.path.resolve() for item in result.measurements}
        overlap = source_paths.intersection(current_paths)
        if overlap:
            raise ValueError(f"Source files assigned to multiple datasets: {sorted(overlap)}")
        source_paths.update(current_paths)

        workbook_path = args.output_root / f"{spec.basename}.xlsx"
        project_path = args.output_root / f"{spec.basename}.json"
        write_workbook(result, workbook_path)
        project = build_project(result, template)
        validate_project(project)
        common.write_project(project, project_path)
        summaries.append(
            {
                "dataset": spec.basename,
                "source_files": len(result.manifest),
                "modes": result.manifest["mode"].value_counts().to_dict(),
                "warnings": int((result.manifest["status"] == "warning").sum()),
                "errors": int((result.manifest["status"] == "error").sum()),
                "workbook": str(workbook_path),
                "project": str(project_path),
                "sheets": len(project["sheets"]),
                "graphs": len(project["graphs"]),
            }
        )

    expected_source_files = {
        path.resolve()
        for sample_id in ("UMD_13", "UMD_14")
        for path in (args.source_root / sample_id).rglob("*.xls")
    }
    if source_paths != expected_source_files:
        missing = sorted(str(path) for path in expected_source_files - source_paths)
        extra = sorted(str(path) for path in source_paths - expected_source_files)
        raise ValueError(f"Campaign coverage mismatch; missing={missing}, extra={extra}")
    if len(source_paths) != 49:
        raise ValueError(f"Expected 49 unique UMD XLS files, found {len(source_paths)}")

    print(json.dumps(summaries, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
